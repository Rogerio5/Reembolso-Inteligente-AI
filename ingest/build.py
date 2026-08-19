"""Constrói os índices da base normativa em storage/.

Executar fora do container:

    python -m ingest.build

O build persiste:
- índice vetorial LlamaIndex;
- índice BM25;
- manifesto dos documentos/chunks.

Nenhuma regra de negócio é codificada aqui.
"""

from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

import fitz
from docx import Document as DocxDocument
from llama_index.core import Document, VectorStoreIndex
from llama_index.core.node_parser import SentenceSplitter
from llama_index.retrievers.bm25 import BM25Retriever

from app.llm import criar_embeddings_llamaindex


RAIZ = Path(__file__).resolve().parents[1]
DIR_KB = RAIZ / "kb"
DIR_STORAGE = RAIZ / "storage"
DIR_VECTOR = DIR_STORAGE / "vector"
DIR_BM25 = DIR_STORAGE / "bm25"
ARQUIVO_MANIFESTO = DIR_STORAGE / "manifest.json"

EXTENSOES = {".pdf", ".docx"}

CHUNK_SIZE = 700
CHUNK_OVERLAP = 100


def _sha256(caminho: Path) -> str:
    digest = hashlib.sha256()

    with caminho.open("rb") as arquivo:
        for bloco in iter(
            lambda: arquivo.read(1024 * 1024),
            b"",
        ):
            digest.update(bloco)

    return digest.hexdigest()


def _carregar_pdf(caminho: Path) -> list[Document]:
    documentos: list[Document] = []

    with fitz.open(caminho) as pdf:
        for numero, pagina in enumerate(pdf, start=1):
            texto = pagina.get_text("text").strip()

            if not texto:
                continue

            documentos.append(
                Document(
                    text=texto,
                    metadata={
                        "arquivo": caminho.name,
                        "pagina": numero,
                        "tipo": "pdf",
                    },
                )
            )

    return documentos


def _carregar_docx(caminho: Path) -> list[Document]:
    docx = DocxDocument(str(caminho))

    texto = "\n".join(
        paragrafo.text.strip()
        for paragrafo in docx.paragraphs
        if paragrafo.text.strip()
    )

    if not texto:
        return []

    return [
        Document(
            text=texto,
            metadata={
                "arquivo": caminho.name,
                "pagina": None,
                "tipo": "docx",
            },
        )
    ]


def _carregar_documentos(
    arquivos: list[Path],
) -> list[Document]:
    documentos: list[Document] = []

    for caminho in arquivos:
        if caminho.suffix.lower() == ".pdf":
            documentos.extend(
                _carregar_pdf(caminho)
            )

        elif caminho.suffix.lower() == ".docx":
            documentos.extend(
                _carregar_docx(caminho)
            )

    return documentos


def _limpar_indices() -> None:
    for caminho in (DIR_VECTOR, DIR_BM25):
        if caminho.exists():
            shutil.rmtree(caminho)

    DIR_STORAGE.mkdir(
        parents=True,
        exist_ok=True,
    )


def main() -> int:
    if not DIR_KB.exists():
        raise FileNotFoundError(
            f"kb/ não encontrado: {DIR_KB}"
        )

    arquivos = sorted(
        caminho
        for caminho in DIR_KB.iterdir()
        if (
            caminho.is_file()
            and caminho.suffix.lower() in EXTENSOES
        )
    )

    if not arquivos:
        raise RuntimeError(
            "Nenhum PDF ou DOCX encontrado em kb/."
        )

    print(
        f"KB_ARQUIVOS={len(arquivos)}"
    )

    documentos = _carregar_documentos(
        arquivos
    )

    if not documentos:
        raise RuntimeError(
            "Nenhum texto foi extraído da kb/."
        )

    print(
        f"KB_DOCUMENTOS_EXTRAIDOS={len(documentos)}"
    )

    splitter = SentenceSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
    )

    nodes = splitter.get_nodes_from_documents(
        documentos
    )

    if not nodes:
        raise RuntimeError(
            "Nenhum chunk foi produzido."
        )

    print(
        f"KB_CHUNKS={len(nodes)}"
    )

    _limpar_indices()

    print(
        "VECTOR_INDEX_BUILD=START"
    )

    embeddings = criar_embeddings_llamaindex()

    indice = VectorStoreIndex(
        nodes,
        embed_model=embeddings,
    )

    indice.storage_context.persist(
        persist_dir=str(DIR_VECTOR)
    )

    print(
        "VECTOR_INDEX_BUILD=OK"
    )

    print(
        "BM25_INDEX_BUILD=START"
    )

    bm25 = BM25Retriever.from_defaults(
        nodes=nodes,
        language="portuguese",
        similarity_top_k=8,
    )

    bm25.persist(
        str(DIR_BM25)
    )

    print(
        "BM25_INDEX_BUILD=OK"
    )

    manifesto = {
        "chunk_size": CHUNK_SIZE,
        "chunk_overlap": CHUNK_OVERLAP,
        "documentos": [
            {
                "arquivo": caminho.name,
                "sha256": _sha256(caminho),
            }
            for caminho in arquivos
        ],
        "quantidade_arquivos": len(arquivos),
        "quantidade_documentos": len(documentos),
        "quantidade_chunks": len(nodes),
        "vector_dir": "vector",
        "bm25_dir": "bm25",
    }

    ARQUIVO_MANIFESTO.write_text(
        json.dumps(
            manifesto,
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )
        + "\n",
        encoding="utf-8",
    )

    print(
        f"MANIFESTO={ARQUIVO_MANIFESTO}"
    )
    print(
        "INGEST_BUILD_OK=True"
    )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
